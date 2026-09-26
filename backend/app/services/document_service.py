# backend/app/services/document_service.py

import os
import csv
import json  # <-- EKLENDİ: Data Lake'e JSON kaydetmek için
import tempfile
import logging
import uuid
from datetime import datetime
from typing import List, Dict, Any, Optional

from pypdf import PdfReader
from docx import Document as DocxDocument
from sqlalchemy.orm import Session

from app.models.models import UploadedDocument
from app.clients.milvus_client import get_milvus_client, COLLECTION_NAME
from app.helpers.text_processor import TextProcessor, text_processor
from app.clients.openai_client import OpenAIClient
from app.clients.rustfs_client import RustFSClient
from app.configs.config import settings  # <-- EKLENDİ: Dinamik bucket isimlerini çekmek için

logger = logging.getLogger("Document_Service")

# TextProcessor instance (Chonkie chunker dahil)
_text_processor = TextProcessor()


# 1. Belge Listeleme (GET /documents)

def list_documents(db: Session) -> List[Dict[str, Any]]:
    """
    Sisteme yüklenmiş olan tüm belgeleri veritabanından çekerek listeler.
    """
    logger.info("Yüklenen belgeler veritabanından listeleniyor.")
    documents = db.query(UploadedDocument).order_by(UploadedDocument.uploaded_at.desc()).all()
    return [
        {
            "id": doc.id,
            "filename": doc.filename,
            "file_type": doc.file_type,
            "file_size": doc.file_size,
            "status": doc.status,
            "uploaded_at": doc.uploaded_at.strftime("%Y-%m-%d %H:%M:%S")
        }
        for doc in documents
    ]


# 2. Belge Silme (DELETE /documents/{doc_id})

def delete_document(doc_id: int, db: Session) -> Optional[Dict[str, str]]:
    """
    Belirtilen belgeyi hem PostgreSQL'den hem RustFS S3 bucket'ından siler hem de
    Milvus vektör veritabanındaki o belgeye ait tüm chunk'ları temizler.
    Belge bulunamazsa None döner.
    """
    doc = db.query(UploadedDocument).filter(UploadedDocument.id == doc_id).first()
    if not doc:
        logger.warning(f"Silinmek istenen belge bulunamadı [ID: {doc_id}]")
        return None

    filename = doc.filename
    logger.info(f"Belge silme süreci başlatıldı: {filename} [ID: {doc_id}]")

    # 1. Milvus'tan bu belgeye ait tüm vektörleri temizliyoruz
    try:
        client = get_milvus_client()
        if client.has_collection(COLLECTION_NAME):
            client.delete(
                collection_name=COLLECTION_NAME,
                filter=f"doc_id == {doc_id}"
            )
            logger.info(f"Milvus vektörleri temizlendi [Doc ID: {doc_id}]")
    except Exception as e:
        logger.error(f"Milvus vektörleri silinirken hata: {e}")

    # 2. RustFS S3 bucket'ından orijinal belge dosyasını temizliyoruz
    try:
        RustFSClient.delete_file(object_name=f"{doc_id}_{filename}")
        # Not: İstenirse JSON yedekleri (processed-chunks) de buradan silinebilir.
        # Şimdilik Data Lake mantığında geriye dönük log olarak tutuyoruz.
        logger.info(f"RustFS S3 deposundan dosya silindi: {doc_id}_{filename}")
    except Exception as e:
        logger.warning(f"[RustFS] Belge dosyası silinirken uyarı: {e}")

    # 3. PostgreSQL'den belge kaydını siliyoruz
    db.delete(doc)
    db.commit()
    logger.info(f"Belge veritabanından başarıyla silindi [ID: {doc_id}]")

    return {
        "status": "success",
        "message": f"'{filename}' başarıyla veritabanından, RustFS S3 deposundan ve Milvus vektör indeksinden silindi."
    }


# 3. Belge Yükleme ve İşleme (POST /upload)

def process_upload(filename: str, file_ext: str, file_content: bytes, db: Session) -> Dict[str, Any]:
    """
    Yüklenen dosyayı işler: orijinal dosyayı doğrudan RustFS S3 bucket'ına yükler,
    metin çıkarır, Chonkie ile parçalar, OpenAI ile embed eder ve Milvus'a yazar.
    """
    file_size = len(file_content)
    logger.info(f"Yeni belge yükleme ve işleme süreci başlatıldı: {filename} ({file_ext})")

    # Veritabanına 'processing' durumunda ön kayıt atıyoruz
    db_doc = UploadedDocument(
        filename=filename,
        file_type=file_ext.replace('.', ''),
        file_size=file_size,
        status="processing",
        uploaded_at=datetime.utcnow()
    )
    db.add(db_doc)
    db.commit()
    db.refresh(db_doc)

    temp_path = None
    try:
        # 1. Orijinal dosyayı DOĞRUDAN RustFS S3 bucket'ına (raw-documents) yüklüyoruz
        object_name = f"{db_doc.id}_{filename}"
        rustfs_object = RustFSClient.upload_file(
            file_bytes=file_content,
            object_name=object_name
        )
        logger.info(f"Dosya RustFS'e yüklendi: {object_name}")

        # Dosyayı geçici bir yere yazma (metin çıkarma kütüphaneleri için)
        with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as temp_file:
            temp_file.write(file_content)
            temp_path = temp_file.name

        processed_chunks = []

        # 2. Dosya Türüne Göre Metin Okuma ve Parçalama
        if file_ext == '.pdf':
            processed_chunks = _extract_pdf_chunks(temp_path)
        elif file_ext == '.docx':
            processed_chunks = _extract_docx_chunks(temp_path)
        elif file_ext == '.csv':
            processed_chunks = _extract_csv_chunks(temp_path)
        elif file_ext == '.txt':
            processed_chunks = _extract_txt_chunks(temp_path)

        if not processed_chunks:
            raise ValueError("Dosyadan anlamlı veya okunabilir metin parçası çıkarılamadı.")

        # --- YENİ EKLENEN KISIM: DATA LAKE JSON YEDEKLEMESİ ---
        _save_chunks_to_datalake(processed_chunks, filename, db_doc.id)
        # ------------------------------------------------------

        # 3. Milvus'a vektörleştirip yazma
        total_chunks = _embed_and_store(processed_chunks, db_doc)

        # Güncelleme: Başarılı
        db_doc.status = "success"
        db.commit()
        logger.info(f"'{filename}' başarıyla işlendi ve indekslendi. Toplam parça: {total_chunks}")

        return {
            "status": "success",
            "doc_id": db_doc.id,
            "filename": filename,
            "file_type": db_doc.file_type,
            "file_size": file_size,
            "chunks_created": total_chunks,
            "rustfs_object": rustfs_object,
            "message": f"'{filename}' başarıyla doğrudan RustFS S3 deposuna yüklendi ve {total_chunks} adet anlamsal parça Milvus'a indekslendi."
        }

    except Exception as e:
        logger.error(f"Belge işlenirken kritik hata oluştu ({filename}): {str(e)}")
        db_doc.status = "error"
        db.commit()
        raise e

    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)


# --- YENİ EKLENEN YARDIMCI FONKSİYON: JSON YEDEKLEME ---

def _save_chunks_to_datalake(chunks: List[Dict[str, Any]], filename: str, doc_id: int):
    """
    Parçalanmış (chunking) metinleri JSON formatında RustFS 'processed-chunks' kovasına yedekler.
    Bu sayede yeni bir embedding modeline geçiş yapılmak istendiğinde PDF'leri baştan okutmaya gerek kalmaz.
    """
    try:
        # 1. Parçaları JSON formatına (UTF-8) çeviriyoruz
        json_bytes = json.dumps(chunks, ensure_ascii=False, indent=2).encode("utf-8")

        # 2. Dosya ismini belirliyoruz (Örn: 22_1955.pdf_chunks.json)
        json_object_name = f"{doc_id}_{filename}_chunks.json"

        # 3. config'den dinamik olarak kova adını çekip yüklüyoruz
        RustFSClient.upload_file(
            file_bytes=json_bytes,
            object_name=json_object_name,
            bucket_name=getattr(settings, "RUSTFS_CHUNKS_BUCKET", "processed-chunks")
        )
        logger.info(f"[RustFS Data Lake] JSON metin parçaları başarıyla yedeklendi: {json_object_name}")
    except Exception as e:
        # Bu yedekleme işlemi kritik değil, asıl süreci (Milvus kaydını) durdurmamak için sadece logluyoruz.
        logger.warning(f"[RustFS Data Lake] JSON yedeklenirken hata oluştu (İşlem devam edecek): {e}")


# --------------------------------------------------------


# Dosya Türüne Göre Metin Çıkarma (Yardımcı)

def _extract_pdf_chunks(temp_path: str) -> List[Dict[str, Any]]:
    """PDF dosyasından hiyerarşik başlık algılama (Madde bazlı) ve uzun maddeler için akıllı alt parçalama uygular."""
    chunks = []
    from pypdf import PdfReader
    reader = PdfReader(temp_path)

    # 1. Sayfa sonu kopmalarını engellemek için tüm metni birleştir
    full_text = ""
    for page in reader.pages:
        full_text += (page.extract_text() or "") + "\n\n"

    # 2. Senin kusursuz Madde ayırıcınla ana blokları bul
    sections = TextProcessor.split_by_madde(full_text, "Yuklenen_Belge")

    for doc in sections:
        section_content = doc.page_content.strip()
        if not section_content:
            continue

        madde_no = doc.metadata.get("madde_no", "Genel_Hukum")
        page_val = madde_no if str(madde_no).isdigit() else 0
        kaynak_val = doc.metadata.get("kaynak", "Yuklenen_Belge")

        # 3. YENİ ve %100 GÜVENLİ: Basit Paragraf (Fıkra) Bölücü
        # Eğer madde 1000 karakterden (yaklaşık 150 kelime) uzunsa, alt parçalara (fıkralara) ayır
        if len(section_content) > 1000:
            paragraphs = section_content.split('\n')
            current_chunk = ""

            for p in paragraphs:
                p = p.strip()
                if not p:
                    continue

                # Eğer mevcut parça ve yeni eklenen paragraf 1000 karakteri aşıyorsa, mevcut parçayı listeye ekle
                if len(current_chunk) + len(p) > 1000 and current_chunk:
                    # Yapay zeka bağlamı kaybetmesin diye başına Madde Numarasını ekliyoruz
                    context_prefix = f"[MADDE {madde_no} - Devamı]\n"
                    chunks.append({
                        "text": context_prefix + current_chunk.strip(),
                        "page": page_val,
                        "madde_no": madde_no,
                        "kaynak": kaynak_val
                    })
                    current_chunk = p  # Yeni parçaya geç
                else:
                    current_chunk += "\n" + p if current_chunk else p

            # Kalan son parçayı (fıkraları) listeye ekle
            if current_chunk:
                context_prefix = f"[MADDE {madde_no} - Devamı]\n"
                chunks.append({
                    "text": context_prefix + current_chunk.strip(),
                    "page": page_val,
                    "madde_no": madde_no,
                    "kaynak": kaynak_val
                })
        else:
            # Madde kısaysa bölmeye gerek yok, doğrudan ekle
            chunks.append({
                "text": section_content,
                "page": page_val,
                "madde_no": madde_no,
                "kaynak": kaynak_val
            })

    return chunks


def _extract_docx_chunks(temp_path: str) -> List[Dict[str, Any]]:
    """Word (DOCX) dosyasından paragraf bloklama ve chunking ile metin çıkarır."""
    chunks = []
    doc = DocxDocument(temp_path)
    current_paragraph_block = []
    word_count = 0

    for p in doc.paragraphs:
        p_text = p.text.strip()
        if not p_text:
            continue
        current_paragraph_block.append(p_text)
        word_count += len(p_text.split())

        if word_count > 300:
            block_text = "\n".join(current_paragraph_block)
            chonkie_chunks = _text_processor.chunk_text(block_text)
            for c in chonkie_chunks:
                chunks.append({"text": c.text.strip(), "page": 1})
            current_paragraph_block = []
            word_count = 0

    if current_paragraph_block:
        block_text = "\n".join(current_paragraph_block)
        chonkie_chunks = _text_processor.chunk_text(block_text)
        for c in chonkie_chunks:
            chunks.append({"text": c.text.strip(), "page": 1})

    return chunks


def _extract_csv_chunks(temp_path: str) -> List[Dict[str, Any]]:
    """CSV dosyasını satır satır okuyup metin parçalarına dönüştürür."""
    chunks = []
    with open(temp_path, mode='r', encoding='utf-8-sig') as csv_file:
        sample = csv_file.read(2048)
        csv_file.seek(0)
        dialect = csv.Sniffer().sniff(sample) if sample else csv.excel

        reader = csv.DictReader(csv_file, dialect=dialect)
        row_idx = 1
        for row in reader:
            row_items = [f"{col}: {val.strip()}" for col, val in row.items() if val and val.strip()]
            if row_items:
                row_text = ", ".join(row_items)
                chunks.append({"text": f"Satır {row_idx}: {row_text}", "page": row_idx})
            row_idx += 1

    return chunks


def _extract_txt_chunks(temp_path: str) -> List[Dict[str, Any]]:
    """TXT dosyasını okuyup Chonkie ile parçalar."""
    chunks = []
    with open(temp_path, mode='r', encoding='utf-8') as txt_file:
        text_content = txt_file.read()
    clean_text = TextProcessor.clean_and_normalize_text(text_content)
    if clean_text:
        chonkie_chunks = _text_processor.chunk_text(clean_text)
        for c in chonkie_chunks:
            chunks.append({"text": c.text.strip(), "page": 1})
    return chunks


# Embedding ve Milvus Yazma (Yardımcı)

def _embed_and_store(processed_chunks: List[Dict], db_doc: UploadedDocument) -> int:
    """Chunk'ları OpenAI ile vektörleştirip Milvus'a yazar. Toplam chunk sayısını döner."""
    client = get_milvus_client()
    if not client.has_collection(COLLECTION_NAME):
        # Koleksiyon yoksa, OpenAI embedding boyutuna (1536) uygun şekilde oluşturuyoruz:
        client.create_collection(
            collection_name=COLLECTION_NAME,
            dimension=1536
        )

    embeddings_model = OpenAIClient.get_embeddings()

    batch_size = 50
    total_chunks = len(processed_chunks)

    for i in range(0, total_chunks, batch_size):
        batch = processed_chunks[i:i + batch_size]
        texts = [c["text"] for c in batch]

        vectors = embeddings_model.embed_documents(texts)

        data = []
        for idx, doc in enumerate(batch):
            data.append({
                "id": uuid.uuid4().int & ((1 << 63) - 1),  # <--- Hatanın çözümü! Milvus'un istediği zorunlu anahtar.
                "text": doc["text"],
                "page": doc["page"],
                "vector": vectors[idx],
                "doc_id": db_doc.id,
                "filename": db_doc.filename
            })

        client.insert(
            collection_name=COLLECTION_NAME,
            data=data
        )

    return total_chunks